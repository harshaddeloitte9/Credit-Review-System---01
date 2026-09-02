"""
Proves the application is genuinely data-driven, not a Tata Steel dashboard
with a file-picker bolted on.

Dataset A = the real reference workbook (Credit Review Automation Example.xlsx),
            unmodified.
Dataset B = a synthetic copy, generated programmatically by this script, with
            Revenue, Total Debt components, the Reported Ratios block (DSCR,
            ISCR, Current Ratio, Debt-Equity, Debt/Assets, margins, Net
            Worth), Collateral book values, the Limit Setting loan amount,
            and Portfolio Monitoring inputs all changed to different,
            non-Tata-Steel numbers.

Both are run through the exact same pipeline app.py uses (the same
build_*/load_* functions, none re-implemented here), and this script
asserts:
  1. Dataset A's results are still pinned to the frozen baseline (regression).
  2. Dataset B's results differ from Dataset A everywhere they logically
     should (Financial Analysis, Risk Scoring, Covenants, Collateral, Limit
     Setting, Continuous Monitoring, Credit Memo).
  3. None of Dataset B's results equal a hardcoded Tata Steel figure -
     i.e. nothing fell back to the reference workbook's own numbers.
  4. workbook_validation.py correctly REJECTS a structurally broken upload
     (missing a hard-required sheet) with a clear message, and correctly
     WARNS (without rejecting) on a workbook missing only the Collateral
     sheet - proving the "unavailable, not a crash, not a Tata Steel
     fallback" behaviour required for modules 5/8.

Dataset B is generated fresh into a temp file on every run; nothing fake is
written into the production application or committed to the repository.
"""

import os
import shutil
import sys
import tempfile

from openpyxl import load_workbook

from reference_data import load_reference_bundle, REFERENCE_WORKBOOK_PATH
from input_data_engine import compute_derived_metrics, run_validation_checks
from financial_analysis import build_financial_analysis
from risk_scoring import build_risk_scoring
from covenants import load_covenant_policy, build_covenants
from collateral_assessment import load_collateral_inputs, build_collateral_assessment
from limit_setting import load_limit_setting_policy, build_limit_setting, LimitSettingResult
from continuous_monitoring import load_monitoring_policy, build_continuous_monitoring
from portfolio_monitoring import load_portfolio_inputs, build_portfolio_monitoring
from credit_memo import load_credit_memo_commentary, build_credit_memo
from workbook_validation import validate_workbook

TATA_STEEL_VALUES = {
    "revenue": 139720.22,
    "total_debt": 69049.26999999999,
    "dscr": 2.85,
    "iscr": 8.08,
    "current_ratio": 0.58,
    "debt_equity": 0.50,
    "risk_score": 3.7875,
}

# (sheet, cell) -> new value. Every cell here is a RAW input in the source
# workbook (confirmed during the earlier architecture audit - none of them
# are formulas), so changing them needs no Excel recalculation: this
# application's own Python layer recomputes everything downstream from
# these raw cells, exactly as it does for the real workbook.
DATASET_B_CHANGES = {
    ("1-Input Data", "B5"): 95000.0,       # Revenue, FY-current (was 139,720.22)
    ("1-Input Data", "C5"): 100000.0,      # Revenue, FY-previous (revenue DECLINED this year)
    ("1-Input Data", "B8"): 78000.0,       # Total expenses, FY-current
    ("1-Input Data", "C8"): 82000.0,       # Total expenses, FY-previous
    ("1-Input Data", "B24"): 180000.0,     # Total assets
    ("1-Input Data", "B25"): 90000.0,      # Total equity
    ("1-Input Data", "B26"): 20000.0,      # Borrowings - non-current (was 50,666.45)
    ("1-Input Data", "B27"): 8000.0,       # Borrowings - current (was 14,037.73)
    ("1-Input Data", "B28"): 1000.0,       # Lease liabilities - non-current
    ("1-Input Data", "B29"): 300.0,        # Lease liabilities - current
    ("1-Input Data", "B49"): 1.10,         # Reported DSCR (was 2.85) - breaches 1.5x covenant
    ("1-Input Data", "B50"): 1.80,         # Reported ISCR (was 8.08) - breaches 2.0x covenant
    ("1-Input Data", "B51"): 1.30,         # Reported Current Ratio (was 0.58) - stronger
    ("1-Input Data", "B52"): 0.15,         # Reported Debt-Equity (was 0.50) - lower leverage
    ("1-Input Data", "B53"): 0.15,         # Reported Debt/Assets (was 0.25)
    ("1-Input Data", "B54"): 0.12,         # Reported EBITDA margin (was 0.2364)
    ("1-Input Data", "B55"): 0.05,         # Reported Net profit margin (was 0.115)
    ("1-Input Data", "B56"): 88000.0,      # Reported Net worth (was 132,345.31)
    ("5-Collateral Assessment", "B15"): 2000.0,    # Cash & Bank Balances (was 3,104)
    ("5-Collateral Assessment", "B16"): 30000.0,   # Investments (was 94,529)
    ("5-Collateral Assessment", "B17"): 1000.0,    # Trade Receivables (was 1,632)
    ("5-Collateral Assessment", "B18"): 8000.0,    # Inventory (was 23,059)
    ("5-Collateral Assessment", "B19"): 60000.0,   # Property, Plant & Equipment (was 109,060)
    ("6-Limit Setting", "B4"): 12000.0,            # Term loan amount requested (was 5,000)
    ("8-Portfolio Monitoring", "C8"): 150000.0,    # Total Portfolio Exposure (was 100,000)
    ("8-Portfolio Monitoring", "C9"): 300,         # Number of Accounts (was 250)
    ("8-Portfolio Monitoring", "C11"): 12000.0,    # Proposed Loan Amount (was 5,000)
    ("8-Portfolio Monitoring", "C13"): 5000.0,     # Gross NPL (was 2,000)
}


def generate_dataset_b(source_path, dest_path):
    wb = load_workbook(source_path, data_only=True)
    for (sheet, cell), value in DATASET_B_CHANGES.items():
        wb[sheet][cell] = value
    wb.save(dest_path)


def generate_broken_workbook_missing_hard_sheet(source_path, dest_path):
    """Deletes '4-Covenants' entirely - a hard-required sheet - to prove
    the validator REJECTS with a clear message rather than letting a raw
    KeyError reach the UI."""
    wb = load_workbook(source_path, data_only=True)
    del wb["4-Covenants"]
    wb.save(dest_path)


def generate_workbook_missing_collateral(source_path, dest_path):
    """Deletes '5-Collateral Assessment' only - everything else intact -
    to prove the validator WARNS (not rejects) and the Collateral module
    degrades to 'unavailable' rather than crashing or reusing Tata Steel's
    collateral numbers."""
    wb = load_workbook(source_path, data_only=True)
    del wb["5-Collateral Assessment"]
    wb.save(dest_path)


# '1-Input Data' rows 49-53 (DSCR/ISCR/Current Ratio/Debt-Equity/Debt-Assets,
# both B and D columns) - the "Reported Ratios" / Regulation 52(4)-style
# disclosure block. Tata Steel's own filing happens to include it, but nothing
# requires a genuinely different borrower's filing to - a real, plausible
# document shape this application must not crash on (see financial_analysis.py /
# risk_scoring.py / covenants.py / continuous_monitoring.py's None-guards).
_REPORTED_RATIOS_CELLS = ["B49", "D49", "B50", "D50", "B51", "D51", "B52", "D52", "B53", "D53"]


def generate_workbook_missing_reported_ratios(source_path, dest_path):
    """Blanks the entire Reported Ratios block - everything else (including
    Debtors/Inventory Turnover Days) intact - to prove Financial Analysis,
    Risk Scoring, Covenants and Continuous Monitoring all degrade to 'Not
    Available / Requires Review' rather than crashing on None arithmetic."""
    wb = load_workbook(source_path, data_only=True)
    ws = wb["1-Input Data"]
    for cell in _REPORTED_RATIOS_CELLS:
        ws[cell] = None
    wb.save(dest_path)


def run_pipeline(workbook_path):
    """Mirrors exactly what app.py does for a given workbook path - no
    calculation logic is re-implemented here, only orchestration."""
    bundle = load_reference_bundle(path=workbook_path)
    d = bundle.input_data
    fa = build_financial_analysis(bundle)
    risk = build_risk_scoring(bundle)

    cov_policy = load_covenant_policy(bundle.workbook_path)
    covenants = build_covenants(fa, cov_policy)

    coll_inputs = load_collateral_inputs(bundle.workbook_path)
    collateral = build_collateral_assessment(coll_inputs, fa.derived.total_debt.current) if coll_inputs else None

    limit_policy = load_limit_setting_policy(bundle.workbook_path)
    # Same EBIT/reported-DSCR availability guard app.py itself applies
    # before calling build_limit_setting() (which has no None-guard of its
    # own) - see AREAS[6] / the Credit Memo page in app.py.
    if fa.derived.ebit.current is None or d.reported_dscr.current is None:
        limit_result = LimitSettingResult(
            annual_instalment=None, schedule=[], implied_existing_debt_service=None,
            total_proforma_debt_service=None, ebit=None, proforma_dscr=None,
            proforma_dscr_status="Not available", max_serviceable_debt_service=None,
            available_additional_debt_service=None, capacity_cap=None,
            recommended_limit=None, conclusion="Not available",
        )
    else:
        limit_result = build_limit_setting(limit_policy, fa.derived.ebit.current, d.reported_dscr.current)

    mon_policy = load_monitoring_policy(bundle.workbook_path)
    rows_by_key = {r.key: r for r in fa.income_statement_rows + fa.ratio_rows}
    breach = any(r.rag in ("AMBER", "RED") for r in covenants.rows)
    monitoring = build_continuous_monitoring(
        mon_policy, rows_by_key["revenue"].pct_change, rows_by_key["total_debt"].pct_change,
        rows_by_key["dscr"].pct_change, breach,
    )

    portfolio_inputs = load_portfolio_inputs(bundle.workbook_path)
    portfolio = build_portfolio_monitoring(portfolio_inputs) if portfolio_inputs else None

    commentary = load_credit_memo_commentary(bundle.workbook_path)
    memo = build_credit_memo(fa, risk, covenants, collateral, limit_result, monitoring, portfolio, commentary)

    return {
        "revenue": d.revenue.current,
        "total_debt": fa.derived.total_debt.current,
        "dscr": d.reported_dscr.current,
        "iscr": d.reported_iscr.current,
        "current_ratio": d.reported_current_ratio.current,
        "debt_equity": d.reported_debt_equity.current,
        "risk_score": risk.total_weighted_score,
        "risk_grade": risk.risk_grade,
        "financial_strength": risk.financial_strength_score,
        "fa_overall_status": fa.overall_status,
        "covenant_overall_status": covenants.overall_status,
        "covenant_dscr_rag": covenants.rows[0].rag,
        "covenant_iscr_rag": covenants.rows[1].rag,
        "collateral_total_adjusted": collateral.total_adjusted_value if collateral else None,
        "collateral_coverage_ratio": collateral.coverage_ratio if collateral else None,
        "collateral_rating": collateral.rating if collateral else None,
        "limit_recommended": limit_result.recommended_limit,
        "limit_conclusion": limit_result.conclusion,
        "limit_annual_instalment": limit_result.annual_instalment,
        "monitoring_overall_status": monitoring.overall_status,
        "monitoring_trigger_statuses": tuple(t.status for t in monitoring.triggers),
        "portfolio_overall_status": portfolio.overall_status if portfolio else None,
        "portfolio_total_exposure": portfolio_inputs.total_portfolio_exposure if portfolio_inputs else None,
        "memo_statuses": tuple(a.status for a in memo.areas),
    }


def _fail(msg):
    print(f"  FAIL: {msg}")
    return False


def _ok(msg):
    print(f"  OK:   {msg}")
    return True


def main():
    all_ok = True
    scratch = tempfile.mkdtemp(prefix="cr_tool_dynamic_test_")
    try:
        dataset_b_path = os.path.join(scratch, "dataset_b_synthetic.xlsx")
        broken_path = os.path.join(scratch, "broken_missing_covenants.xlsx")
        no_collateral_path = os.path.join(scratch, "missing_collateral_only.xlsx")
        no_reported_ratios_path = os.path.join(scratch, "missing_reported_ratios.xlsx")

        print("=" * 100)
        print("GENERATING TEST WORKBOOKS (temp files, not committed to the repo)")
        print("=" * 100)
        generate_dataset_b(REFERENCE_WORKBOOK_PATH, dataset_b_path)
        print(f"  Dataset B written to {dataset_b_path}")
        generate_broken_workbook_missing_hard_sheet(REFERENCE_WORKBOOK_PATH, broken_path)
        print(f"  Broken (missing '4-Covenants') workbook written to {broken_path}")
        generate_workbook_missing_collateral(REFERENCE_WORKBOOK_PATH, no_collateral_path)
        print(f"  Missing-collateral-only workbook written to {no_collateral_path}")
        generate_workbook_missing_reported_ratios(REFERENCE_WORKBOOK_PATH, no_reported_ratios_path)
        print(f"  Missing-Reported-Ratios workbook written to {no_reported_ratios_path}")

        print()
        print("=" * 100)
        print("1. RUNNING FULL PIPELINE ON DATASET A (Tata Steel, unmodified)")
        print("=" * 100)
        result_a = run_pipeline(REFERENCE_WORKBOOK_PATH)
        for k, v in result_a.items():
            print(f"  {k:28s} = {v}")

        bundle_a = load_reference_bundle(path=REFERENCE_WORKBOOK_PATH)
        checks_a = run_validation_checks(bundle_a.input_data, compute_derived_metrics(bundle_a.input_data))
        passed_a = sum(1 for c in checks_a if c.result in ("PASS", "PASS - ties"))
        print(f"  {'input_validation_checks':28s} = {passed_a}/{len(checks_a)}")

        print()
        print("=" * 100)
        print("2. REGRESSION CHECK - Dataset A must still match the frozen baseline")
        print("=" * 100)
        checks = [
            ("risk_score", 3.7875, 1e-6),
            ("risk_grade", "Low Risk", None),
            ("financial_strength", 4.15, 1e-6),
            ("dscr", 2.85, 1e-6),
            ("iscr", 8.08, 1e-6),
            ("current_ratio", 0.58, 1e-6),
            ("debt_equity", 0.50, 1e-6),
        ]
        for key, expected, tol in checks:
            actual = result_a[key]
            if tol is None:
                ok = actual == expected
            else:
                ok = abs(actual - expected) <= tol
            all_ok &= (_ok if ok else _fail)(f"Dataset A {key} = {actual} (expected {expected})")
        all_ok &= (_ok if passed_a == len(checks_a) else _fail)(
            f"Dataset A input validation checks = {passed_a}/{len(checks_a)} (expected {len(checks_a)}/{len(checks_a)})"
        )

        print()
        print("=" * 100)
        print("3. RUNNING FULL PIPELINE ON DATASET B (synthetic, different company)")
        print("=" * 100)
        result_b = run_pipeline(dataset_b_path)
        for k, v in result_b.items():
            if k != "checks_9_9":
                print(f"  {k:28s} = {v}")

        print()
        print("=" * 100)
        print("4. DATASET A != DATASET B - propagation proof")
        print("=" * 100)
        propagation_checks = [
            ("revenue", "Financial Analysis"),
            ("total_debt", "Financial Analysis"),
            ("dscr", "Financial Analysis / Covenants / Risk Scoring / Limit Setting"),
            ("iscr", "Financial Analysis / Covenants / Risk Scoring"),
            ("current_ratio", "Financial Analysis / Risk Scoring"),
            ("debt_equity", "Financial Analysis / Risk Scoring"),
            ("risk_score", "Risk Scoring"),
            ("risk_grade", "Risk Scoring"),
            ("financial_strength", "Risk Scoring"),
            ("fa_overall_status", "Financial Analysis"),
            ("covenant_overall_status", "Covenants"),
            ("covenant_dscr_rag", "Covenants"),
            ("covenant_iscr_rag", "Covenants"),
            ("collateral_total_adjusted", "Collateral Assessment"),
            ("collateral_coverage_ratio", "Collateral Assessment"),
            ("limit_recommended", "Limit Setting"),
            ("limit_annual_instalment", "Limit Setting"),
            ("monitoring_trigger_statuses", "Continuous Monitoring"),
            ("portfolio_total_exposure", "Portfolio Monitoring (own inputs, not borrower financials)"),
            ("memo_statuses", "Credit Memo"),
        ]
        for key, module in propagation_checks:
            a_val, b_val = result_a[key], result_b[key]
            changed = a_val != b_val
            all_ok &= (_ok if changed else _fail)(
                f"{module}: {key} changed ({a_val} -> {b_val})" if changed
                else f"{module}: {key} DID NOT CHANGE ({a_val}) - expected it to differ"
            )

        print()
        print("=" * 100)
        print("5. NO TATA STEEL FALLBACK - Dataset B must not equal any Tata Steel figure")
        print("=" * 100)
        for key, tata_value in TATA_STEEL_VALUES.items():
            b_val = result_b[key]
            if isinstance(tata_value, float):
                distinct = abs(b_val - tata_value) > 1e-6
            else:
                distinct = b_val != tata_value
            all_ok &= (_ok if distinct else _fail)(
                f"Dataset B {key} = {b_val} (Tata Steel value was {tata_value}) - "
                f"{'no fallback detected' if distinct else 'MATCHES TATA STEEL - possible silent fallback'}"
            )

        print()
        print("=" * 100)
        print("6. VALIDATOR: rejects a workbook missing a hard-required sheet")
        print("=" * 100)
        report_broken = validate_workbook(broken_path)
        all_ok &= (_ok if not report_broken.is_valid else _fail)(
            f"Broken workbook (missing '4-Covenants') rejected: is_valid={report_broken.is_valid}, "
            f"errors={report_broken.errors}"
        )

        print()
        print("=" * 100)
        print("7. VALIDATOR: warns (does not reject) on a workbook missing only Collateral")
        print("=" * 100)
        report_no_collateral = validate_workbook(no_collateral_path)
        all_ok &= (_ok if report_no_collateral.is_valid else _fail)(
            f"Missing-collateral-only workbook still valid: is_valid={report_no_collateral.is_valid}"
        )
        all_ok &= (_ok if report_no_collateral.warnings else _fail)(
            f"Missing-collateral-only workbook produced a warning: {report_no_collateral.warnings}"
        )

        print()
        print("=" * 100)
        print("8. GRACEFUL DEGRADATION: Collateral module shows 'unavailable', no crash, no fallback")
        print("=" * 100)
        result_no_collateral = run_pipeline(no_collateral_path)
        all_ok &= (_ok if result_no_collateral["collateral_rating"] is None else _fail)(
            f"Collateral rating for the missing-collateral workbook = {result_no_collateral['collateral_rating']} "
            f"(expected None / unavailable, not a crash and not Tata Steel's rating)"
        )

        print()
        print("=" * 100)
        print("9. VALIDATOR: the real Dataset A and Dataset B both pass cleanly")
        print("=" * 100)
        report_a = validate_workbook(REFERENCE_WORKBOOK_PATH)
        report_b = validate_workbook(dataset_b_path)
        all_ok &= (_ok if report_a.is_valid else _fail)(f"Dataset A valid: {report_a.is_valid}, errors={report_a.errors}")
        all_ok &= (_ok if report_b.is_valid else _fail)(f"Dataset B valid: {report_b.is_valid}, errors={report_b.errors}")

        print()
        print("=" * 100)
        print("10. GRACEFUL DEGRADATION: a document with no Reported Ratios disclosure at all")
        print("    (DSCR/ISCR/Current Ratio/Debt-Equity/Debt-Assets) never crashes, never")
        print("    fabricates a value, never reuses Tata Steel's own reported ratios")
        print("=" * 100)
        result_no_ratios = run_pipeline(no_reported_ratios_path)
        for k in ("dscr", "iscr", "current_ratio", "debt_equity", "risk_score", "risk_grade",
                  "covenant_overall_status", "covenant_dscr_rag", "covenant_iscr_rag"):
            all_ok &= (_ok if result_no_ratios[k] is None or "Not Available" in str(result_no_ratios[k]) else _fail)(
                f"{k} for the missing-Reported-Ratios document = {result_no_ratios[k]!r} "
                f"(expected None / 'Not Available...', not a crash and not a Tata Steel figure)"
            )
        # Corrected expectation (Phase 1 extraction/completeness audit,
        # financial_analysis.py's build_financial_analysis): DSCR, ISCR,
        # Current Ratio and Debt-Equity are each one of the metrics this
        # application treats as REQUIRED for the Approve/Reject decision
        # (see CRITICAL_WEIGHT / _mark() there) - a document missing all
        # four must never resolve to Approve or Reject just because
        # Revenue/Total Assets/Net Profit/OCF still happen to be present.
        # This used to assert the opposite (Approve/Reject "still gate
        # it") - that was the false-completeness bug this audit exists to
        # fix, not a real requirement, so the assertion is corrected here
        # rather than preserved.
        all_ok &= (_ok if result_no_ratios["memo_statuses"][0] == "Not Available / Requires Review" else _fail)(
            f"Financial Analysis overall status is 'Not Available / Requires Review' for the "
            f"missing-Reported-Ratios document ({result_no_ratios['memo_statuses'][0]}) - DSCR/ISCR/"
            f"Current Ratio/Debt-Equity are all required for the decision, so Approve/Reject is "
            f"withheld even though Revenue/Total Assets/Net Profit/OCF are present"
        )

        print()
        print("#" * 100)
        print("OVERALL RESULT:", "ALL CHECKS PASSED - the application is genuinely data-driven" if all_ok
              else "SOME CHECKS FAILED - see FAIL lines above")
        print("#" * 100)
        return 0 if all_ok else 1
    finally:
        shutil.rmtree(scratch, ignore_errors=True)


if __name__ == "__main__":
    sys.exit(main())
