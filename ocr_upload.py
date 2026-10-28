"""
PDF-upload integration layer: wires the validated ocr_pipeline.py functions
into the live Streamlit app.

Nothing here re-implements OCR/extraction logic - it only orchestrates the
existing, validated ocr_pipeline.py functions (run_llmwhisperer,
extract_financial_values, populate_workbook) plus demo_source's own
build_demo_workbook(), exactly as demo_source.py's own docstring anticipated:
"whatever produces a 'Generated Extracted Sheet' from a user's uploaded PDF
can call build_demo_workbook(that_path, POLICY_WORKBOOK_PATH) with no change
needed here or in any downstream module - only the first argument changes."

The one genuinely new piece is the network-transport fallback: this
environment's Python `requests` library gets its connection reset uploading
a large PDF to the LLMWhisperer endpoint (confirmed during OCR-pipeline
validation - WinError 10053, reproducible, unrelated to file content), while
curl reaches the same endpoint with the same file without issue. This tries
the plain SDK path first (so a deployment where that quirk doesn't exist
just works normally) and only falls back to the curl-based transport on a
connection-level failure - never on an HTTP-level error, which is a real
API response and should propagate as-is.
"""

import csv
import hashlib
import json
import os
import subprocess
import tempfile
import time
import urllib.parse
from datetime import datetime

import requests
from openpyxl import load_workbook

import ocr_pipeline
import ai_validation
import ai_correction_proposals
import ai_dynamic_mapping
from demo_source import build_demo_workbook, DemoSourceError

UPLOAD_WORK_DIR = os.path.join(tempfile.gettempdir(), "credit_review_tool_pdf_upload")

CLIENT_API_TIMEOUT = 300


class OcrUploadError(RuntimeError):
    """Raised when OCR, extraction, or workbook generation fails for an
    uploaded PDF. Never silently worked around - surfaced to the caller
    with a clear, specific message."""


def _ensure_work_dir():
    os.makedirs(UPLOAD_WORK_DIR, exist_ok=True)


def _pdf_content_hash(pdf_path):
    """SHA-256 of the PDF's own bytes (not its filename) - two uploads of
    the exact same document hash identically regardless of what the file
    is named, so a re-upload can be recognized as "already OCR'd" even if
    renamed."""
    digest = hashlib.sha256()
    with open(pdf_path, "rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _cached_ocr_path(pdf_hash):
    return os.path.join(UPLOAD_WORK_DIR, f"raw_ocr_cache_{pdf_hash}.txt")


def get_cached_ocr_text(pdf_path):
    """Returns the previously-saved OCR text for this exact PDF content,
    or None if no cached result exists yet - never calls the OCR provider.
    Exposed for callers (e.g. the Data Upload UI) that want to know
    up front whether an upload will need a fresh OCR call before
    triggering process_uploaded_pdf()."""
    _ensure_work_dir()
    cache_path = _cached_ocr_path(_pdf_content_hash(pdf_path))
    if os.path.exists(cache_path):
        with open(cache_path, "r", encoding="utf-8") as f:
            text = f.read()
        if text:
            return text
    return None


def _run_llmwhisperer_via_curl(pdf_path):
    """Exact transport-only fallback validated during the OCR pipeline
    audit: every parameter, the auth header, and the async status-polling/
    retrieve loop are an unmodified replication of what
    LLMWhispererClientV2.whisper() itself sends and does - only the large
    upload POST is sent via curl instead of `requests`."""

    params = {
        "mode": ocr_pipeline.OCR_MODE,
        "output_mode": ocr_pipeline.OCR_OUTPUT_MODE,
        "page_separator": "<<<",
        "pages_to_extract": "",
        "median_filter_size": 0,
        "gaussian_blur_radius": 0,
        "line_splitter_tolerance": 0.4,
        "horizontal_stretch_factor": 1.0,
        "mark_vertical_lines": False,
        "mark_horizontal_lines": False,
        "line_splitter_strategy": "left-priority",
        "add_line_nos": True,
        "include_line_confidence": False,
        "word_confidence_threshold": 0.3,
        "lang": ocr_pipeline.OCR_LANGUAGE,
        "tag": "default",
        "file_name": "",
        "webhook_metadata": "",
        "use_webhook": "",
    }
    query = urllib.parse.urlencode(params)
    url = f"{ocr_pipeline.LLMWHISPERER_BASE_URL}/whisper?{query}"

    _ensure_work_dir()
    response_path = os.path.join(UPLOAD_WORK_DIR, "llmwhisperer_upload_response.json")

    curl_cmd = [
        "curl", "-s",
        "-X", "POST",
        "-H", f"unstract-key: {ocr_pipeline.LLMWHISPERER_API_KEY}",
        "-H", "Content-Type:",
        "--data-binary", f"@{pdf_path}",
        "-o", response_path,
        "-w", "%{http_code}",
        "--max-time", "300",
        url,
    ]

    proc = subprocess.run(curl_cmd, capture_output=True, text=True)
    if proc.returncode != 0:
        raise OcrUploadError(f"Upload transport (curl) failed (exit {proc.returncode}): {proc.stderr}")

    http_code = proc.stdout.strip()
    status_code = int(http_code) if http_code.isdigit() else -1

    with open(response_path, "r", encoding="utf-8") as f:
        body_text = f.read()

    try:
        message = json.loads(body_text) if body_text else {}
        if not isinstance(message, dict):
            message = {"message": str(message)}
    except (json.JSONDecodeError, ValueError):
        message = {"message": body_text}

    message["status_code"] = status_code

    if status_code not in (200, 202):
        message["extraction"] = {}
        raise OcrUploadError(f"LLMWhisperer upload rejected: HTTP {status_code} - {message}")

    client = ocr_pipeline.LLMWhispererClientV2(
        base_url=ocr_pipeline.LLMWHISPERER_BASE_URL,
        api_key=ocr_pipeline.LLMWHISPERER_API_KEY,
        logging_level="ERROR",
    )

    if status_code == 202:
        message["extraction"] = {}
        whisper_hash = message["whisper_hash"]
        start_time = time.time()
        wait_timeout = ocr_pipeline.OCR_WAIT_TIMEOUT

        while time.time() - start_time < wait_timeout:
            status = client.whisper_status(whisper_hash=whisper_hash)
            if status["status_code"] != 200:
                raise OcrUploadError(f"Whisper status check failed: {status}")

            if status["status"] in ("accepted", "processing"):
                pass
            elif status["status"] == "processed":
                resultx = client.whisper_retrieve(whisper_hash=whisper_hash)
                if resultx["status_code"] == 200:
                    message["status_code"] = 200
                    message["extraction"] = resultx["extraction"]
                    break
                raise OcrUploadError(f"Whisper retrieve failed: {resultx}")
            elif status["status"] == "error" or "error" in status["status"]:
                raise OcrUploadError(f"Whisper processing failed: {status}")

            time.sleep(5)
        else:
            raise OcrUploadError("Whisper client operation timed out waiting for completion.")

    extraction = message.get("extraction", {})
    if isinstance(extraction, dict):
        text = extraction.get("result_text", "")
    else:
        text = str(extraction)

    if not text:
        raise OcrUploadError("LLMWhisperer completed, but no extracted text was returned.")

    return text, message


def run_ocr(pdf_path):
    """Run OCR on `pdf_path` via ocr_pipeline.run_llmwhisperer() (unmodified)
    first; only on a connection-level transport failure, retry via the
    curl-based fallback validated for this environment. An HTTP-level
    error (bad API key, quota exceeded, etc.) is a real response and is
    never masked by the fallback - it propagates immediately."""

    try:
        return ocr_pipeline.run_llmwhisperer(pdf_path)
    except (requests.exceptions.ConnectionError, ConnectionAbortedError, ConnectionResetError) as exc:
        try:
            return _run_llmwhisperer_via_curl(pdf_path)
        except Exception as fallback_exc:
            raise OcrUploadError(
                f"OCR upload failed via both the standard transport ({exc}) "
                f"and the curl fallback ({fallback_exc})."
            ) from fallback_exc


# A1's company/scope/unit header text must stay parseable by
# reference_data._parse_header() ("Area 1 - Extracted & Validated Data:
# {company}, {scope} ({unit})"). The template's own default names a
# specific company - replaced here with a neutral placeholder for every
# freshly generated workbook, since the real company name is not
# separately extracted from the OCR text in this build (see the
# integration report's known-limitations section). Never left as the
# template's original name, which would misattribute a different
# uploaded company's data.
GENERIC_BORROWER_LABEL = "Uploaded Borrower"


def _anonymize_header(workbook_path):
    wb = load_workbook(workbook_path)
    ws = wb[ocr_pipeline.TARGET_SHEET]
    ws["A1"] = f"Area 1 - Extracted & Validated Data: {GENERIC_BORROWER_LABEL}, Standalone ({chr(8377)} crore)"
    wb.save(workbook_path)


def _notify(status_callback, stage, state, detail=None):
    if status_callback is not None:
        status_callback(stage, state, detail)


# LLMWhisperer is called with page_separator="<<<" (see
# _run_llmwhisperer_via_curl): every page of the OCR text ends with a line
# that is exactly this separator, and normalize_text() keeps that line, so
# it stays in the same filtered line list every source_line_number indexes.
OCR_PAGE_SEPARATOR = "<<<"


def line_page_numbers(lines):
    """1-based PDF page number for each entry of `lines` (the filtered,
    normalized OCR lines that source_line_number / proposal / mapping line
    numbers all index into). A separator line belongs to the page it ends.
    Returns None when the text has no page separators at all, so a caller
    shows "Not available" rather than guessing a page."""
    if not any(line.strip() == OCR_PAGE_SEPARATOR for line in lines):
        return None
    pages, page = [], 1
    for line in lines:
        pages.append(page)
        if line.strip() == OCR_PAGE_SEPARATOR:
            page += 1
    return pages


def _result_reason(result):
    return getattr(result, "reason", None) or getattr(result, "explanation", None) or ""


def ai_stage_outcome(results, availability, not_needed_detail, unexpected_error=None, report_partial=False):
    """(state, detail) for an advisory AI stage, decided from what actually
    happened - never from the provider setting alone:
      "not_needed"     - nothing needed an AI attempt; the provider was never checked
      "unavailable"    - the provider was checked and is not available
      "exception"      - the stage itself raised unexpectedly
      "attempt_failed" - AI calls were made, but none produced a usable result
      "partial"        - (report_partial only) some fields got a usable AI
                         result, others hit an error or were skipped
      "skipped"        - no AI call was made (time budget, or nothing to send)
      "complete"       - at least one field got a usable result from a real AI call

    Only a result produced by an actual AI call counts as success. A
    result marked ai_call_made=False (e.g. dynamic mapping's
    INSUFFICIENT_CONTEXT when a field has no candidate rows) never does."""
    if unexpected_error is not None:
        return "exception", unexpected_error
    if availability is None or not results:
        return "not_needed", not_needed_detail
    if not availability.reachable or not availability.model_installed:
        return "unavailable", availability.detail
    values = list(results.values())
    total = len(values)
    succeeded = [
        r for r in values
        if r.technical_status == ai_validation.AI_COMPLETED and getattr(r, "ai_call_made", True)
    ]
    no_call = [
        r for r in values
        if r.technical_status == ai_validation.AI_COMPLETED and not getattr(r, "ai_call_made", True)
    ]
    errors = [r for r in values if r.technical_status == ai_validation.AI_ERROR]
    skipped = [r for r in values if r.technical_status == ai_validation.AI_SKIPPED]

    notes = []
    if errors:
        notes.append(f"{len(errors)} of {total} field(s) hit a technical error")
    if skipped:
        notes.append(f"{len(skipped)} of {total} field(s) skipped (time budget exhausted)")
    if no_call:
        notes.append(f"{len(no_call)} of {total} field(s) had no candidate rows to send (no AI call made)")

    if succeeded:
        state = "partial" if report_partial and (errors or skipped) else "complete"
        return state, "; ".join(notes) or None
    if errors:
        return "attempt_failed", "; ".join([_result_reason(errors[0])] + notes)
    if skipped:
        return "skipped", "; ".join([_result_reason(skipped[0])] + notes[1:])
    if no_call:
        return "skipped", "; ".join(notes)
    return "attempt_failed", "No field received a usable AI result."


def process_uploaded_pdf(pdf_path, original_filename, status_callback=None):
    """End-to-end: PDF -> OCR -> extraction -> fresh workbook (from the
    unmodified template) -> merged with the policy workbook. Returns a
    dict with every artifact the UI needs. Raises OcrUploadError /
    ocr_pipeline / demo_source's own exceptions on failure - nothing here
    invents a fallback value.

    status_callback(stage, state, detail=None), if given, is called at
    each real step boundary below (stage in "ocr"/"extraction"/
    "population"; state in "processing"/"complete"/"reused"/"failed"; the
    advisory AI stages may also report "not_needed"/"unavailable"/
    "exception"/"attempt_failed"/"partial"/"skipped" - see
    ai_stage_outcome()) so
    a caller (the Data Upload UI) can show genuine per-stage progress
    rather than a simulated one. Optional and side-effect-free when
    omitted - no behavior here depends on it.

    If OCR has already been run for this exact PDF content (by SHA-256,
    not filename) in a previous call, the cached result is reused and the
    OCR provider is not called again - see _pdf_content_hash /
    get_cached_ocr_text. Extraction and workbook population always run
    fresh from that text regardless, since they are local/deterministic
    and not the quota-limited step."""

    _ensure_work_dir()
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")

    pdf_hash = _pdf_content_hash(pdf_path)
    cache_path = _cached_ocr_path(pdf_hash)
    ocr_text = None
    ocr_reused = False
    if os.path.exists(cache_path):
        with open(cache_path, "r", encoding="utf-8") as f:
            cached_text = f.read()
        if cached_text:
            ocr_text = cached_text
            ocr_reused = True

    if ocr_reused:
        _notify(status_callback, "ocr", "reused")
    else:
        _notify(status_callback, "ocr", "processing")
        try:
            ocr_text, raw_result = run_ocr(pdf_path)
        except Exception as exc:
            _notify(status_callback, "ocr", "failed", str(exc))
            raise
        with open(cache_path, "w", encoding="utf-8") as f:
            f.write(ocr_text)
        _notify(status_callback, "ocr", "complete")

    # Kept as a per-upload, timestamp-named artifact (the existing "Raw
    # OCR Output" download button reads this path) regardless of whether
    # this call hit the content-hash cache above - only the OCR provider
    # call itself is ever skipped, not this artifact.
    raw_txt_path = os.path.join(UPLOAD_WORK_DIR, f"raw_ocr_{timestamp}.txt")
    with open(raw_txt_path, "w", encoding="utf-8") as f:
        f.write(ocr_text)

    _notify(status_callback, "extraction", "processing")
    try:
        extracted, debug_rows, meta = ocr_pipeline.extract_financial_values(
            ocr_text, ocr_pipeline.TEMPLATE_FILE
        )
    except Exception as exc:
        _notify(status_callback, "extraction", "failed", str(exc))
        raise
    _notify(status_callback, "extraction", "complete")

    # `lines_for_ai` (whitespace-collapsed) and `layout_lines_for_ai`
    # (column-position PRESERVING - see ocr_pipeline._normalize_preserve_
    # layout) are reconstructed from `ocr_text` using ocr_pipeline's own
    # (unmodified) helpers, built from the SAME filtered raw/normalized
    # pairs so the two stay index-aligned - exactly the same two lists
    # extract_financial_values() itself builds internally, so
    # source_line_number in `extracted` indexes into both correctly.
    # Kept as a local reconstruction here rather than a new return value
    # from extract_financial_values(), so ocr_pipeline.py's extraction
    # logic and its function signature are untouched by every AI phase.
    # Built HERE (before population) rather than after, so the AI mapping
    # step below can use them to try resolving NOT_FOUND fields before
    # the workbook is ever populated.
    raw_and_normalized = [
        (raw, ocr_pipeline.normalize_text(raw)) for raw in ocr_text.splitlines()
    ]
    raw_and_normalized = [(raw, norm) for raw, norm in raw_and_normalized if norm]
    lines_for_ai = [norm for _raw, norm in raw_and_normalized]
    layout_lines_for_ai = [
        ocr_pipeline._normalize_preserve_layout(raw) for raw, _norm in raw_and_normalized
    ]
    # Display-only provenance: the PDF page of every line above (index-
    # aligned with lines_for_ai), so the Data Review UI can show the page
    # a value came from. Never read by extraction or any calculation.
    source_line_pages = line_page_numbers(lines_for_ai)

    # ------------------------------------------------------------
    # AI-ASSISTED DYNAMIC FIELD MAPPING (this phase) - advisory only,
    # see ai_dynamic_mapping.py. Runs BEFORE population so a verified
    # AI-mapped value is written into the template exactly like any
    # other extracted field - populate_workbook() itself is never
    # modified or aware this phase exists. Same defensive try/except
    # posture as AI validation below: optional, must never break the
    # primary upload the analyst is waiting on. Only ever touches
    # fields ocr_pipeline already marked NOT_FOUND - every other field
    # passes through this call byte-for-byte unchanged.
    # ------------------------------------------------------------
    _notify(status_callback, "ai_mapping", "processing")
    ai_mapping_error = None
    try:
        extracted, ai_mapping_results, ai_mapping_availability = ai_dynamic_mapping.map_missing_fields(
            extracted, lines_for_ai, meta, layout_lines=layout_lines_for_ai
        )
    except Exception as exc:
        ai_mapping_results = {}
        ai_mapping_error = f"AI dynamic mapping failed unexpectedly: {exc}"
        ai_mapping_availability = ai_validation.OllamaAvailability(
            reachable=False, model_installed=False, installed_models=[],
            detail=ai_mapping_error,
        )

    ai_mapping_technical_counts = {}
    ai_mapping_status_counts = {}
    for result in ai_mapping_results.values():
        ai_mapping_technical_counts[result.technical_status] = ai_mapping_technical_counts.get(result.technical_status, 0) + 1
        if result.mapping_status is not None:
            ai_mapping_status_counts[result.mapping_status] = ai_mapping_status_counts.get(result.mapping_status, 0) + 1

    ai_mapping_available = (
        bool(ai_mapping_results)
        and ai_mapping_availability is not None
        and ai_mapping_availability.reachable and ai_mapping_availability.model_installed
    )
    ai_mapping_stage_state, ai_mapping_stage_detail = ai_stage_outcome(
        ai_mapping_results, ai_mapping_availability,
        "No fields required a mapping attempt.", unexpected_error=ai_mapping_error,
        report_partial=True,
    )
    _notify(status_callback, "ai_mapping", ai_mapping_stage_state, ai_mapping_stage_detail)

    _notify(status_callback, "population", "processing")
    try:
        wb = ocr_pipeline.populate_workbook(ocr_pipeline.TEMPLATE_FILE, extracted)
        extracted_xlsx_path = os.path.join(
            UPLOAD_WORK_DIR, f"extracted_{timestamp}.xlsx"
        )
        wb.save(extracted_xlsx_path)
        _anonymize_header(extracted_xlsx_path)

        try:
            merged_path = build_demo_workbook(extracted_sheet_path=extracted_xlsx_path)
        except DemoSourceError as exc:
            raise OcrUploadError(f"Could not merge extracted data with the policy workbook: {exc}") from exc
    except Exception as exc:
        _notify(status_callback, "population", "failed", str(exc))
        raise
    _notify(status_callback, "population", "complete")

    status_counts = {
        "CORRECT": 0, "NOT_FOUND_IN_OCR": 0, "AMBIGUOUS": 0, "NORMALIZED_FROM_OCR": 0,
    }
    for row in debug_rows:
        field = row["Field"]
        ocr_source_line = row["OCR Source Line"]
        fy_current = row["FY2026 Selected"]
        normalized = row["Normalized"]
        ambiguous = row["Ambiguous"]
        if ocr_source_line == "NOT FOUND" or fy_current is None:
            status_counts["NOT_FOUND_IN_OCR"] += 1
        elif ambiguous:
            status_counts["AMBIGUOUS"] += 1
        elif normalized:
            status_counts["NORMALIZED_FROM_OCR"] += 1
        else:
            status_counts["CORRECT"] += 1

    # ------------------------------------------------------------
    # AI-ASSISTED VALIDATION (Phase 2) - advisory only, see ai_validation.py.
    #
    # Deliberately NOT wrapped in the same "raise on failure" pattern as
    # ocr/extraction/population above: this stage is optional and must
    # never break the primary pipeline. ai_validation.validate_extracted_
    # fields() already never raises for an unreachable Ollama or a
    # missing model (both degrade to a clear AI_UNAVAILABLE result per
    # field) - the try/except here is a second line of defense against a
    # genuinely unexpected failure in this module, which still must not
    # fail the upload the analyst is waiting on.
    #
    # `lines_for_ai`/`layout_lines_for_ai` were already built above
    # (before population, for the AI mapping step) - reused here as-is,
    # not rebuilt.
    # ------------------------------------------------------------
    _notify(status_callback, "ai_validation", "processing")
    ai_validation_error = None
    try:
        ai_results, ai_presence_results, ai_availability = ai_validation.validate_extracted_fields(
            extracted, lines_for_ai, meta, layout_lines=layout_lines_for_ai
        )
    except Exception as exc:
        ai_results = {}
        ai_presence_results = {}
        ai_validation_error = f"AI validation failed unexpectedly: {exc}"
        ai_availability = ai_validation.OllamaAvailability(
            reachable=False, model_installed=False, installed_models=[],
            detail=ai_validation_error,
        )

    ai_technical_counts = {}
    ai_validation_counts = {}
    for result in ai_results.values():
        ai_technical_counts[result.technical_status] = ai_technical_counts.get(result.technical_status, 0) + 1
        if result.validation_status is not None:
            ai_validation_counts[result.validation_status] = ai_validation_counts.get(result.validation_status, 0) + 1

    ai_presence_counts = {}
    for result in ai_presence_results.values():
        ai_technical_counts[result.technical_status] = ai_technical_counts.get(result.technical_status, 0) + 1
        if result.presence_status is not None:
            ai_presence_counts[result.presence_status] = ai_presence_counts.get(result.presence_status, 0) + 1

    ai_validation_available = (
        bool(ai_results or ai_presence_results)
        and ai_availability.reachable and ai_availability.model_installed
    )
    _notify(status_callback, "ai_validation", *ai_stage_outcome(
        {**ai_presence_results, **ai_results}, ai_availability,
        "No fields required AI validation.", unexpected_error=ai_validation_error,
    ))

    # ------------------------------------------------------------
    # AI-ASSISTED MAPPING CORRECTION PROPOSALS (Phase 3A) - advisory
    # only, see ai_correction_proposals.py. Runs AFTER AI validation
    # above, since it only considers fields that validation itself (or
    # the deterministic pipeline) already flagged as suspect. Same
    # defensive try/except posture as AI validation: optional, must
    # never break the primary upload the analyst is waiting on.
    # ------------------------------------------------------------
    _notify(status_callback, "ai_correction", "processing")
    correction_error = None
    try:
        correction_proposals, correction_availability = ai_correction_proposals.generate_correction_proposals(
            extracted, ai_results, lines_for_ai, meta, layout_lines=layout_lines_for_ai
        )
    except Exception as exc:
        correction_proposals = {}
        correction_error = f"AI correction-proposal generation failed unexpectedly: {exc}"
        correction_availability = ai_validation.OllamaAvailability(
            reachable=False, model_installed=False, installed_models=[],
            detail=correction_error,
        )

    correction_technical_counts = {}
    correction_proposal_counts = {}
    for proposal in correction_proposals.values():
        correction_technical_counts[proposal.technical_status] = correction_technical_counts.get(proposal.technical_status, 0) + 1
        if proposal.proposal_status is not None:
            correction_proposal_counts[proposal.proposal_status] = correction_proposal_counts.get(proposal.proposal_status, 0) + 1

    if correction_availability is None:
        # No field needed a correction-proposal attempt - Ollama was
        # never even probed (see generate_correction_proposals).
        correction_detail = "No fields required a correction-proposal attempt."
        correction_available = False
    else:
        correction_detail = correction_availability.detail
        correction_available = (
            bool(correction_proposals)
            and correction_availability.reachable and correction_availability.model_installed
        )
    _notify(status_callback, "ai_correction", *ai_stage_outcome(
        correction_proposals, correction_availability,
        "No fields required a correction-proposal attempt.", unexpected_error=correction_error,
    ))

    return {
        "pdf_name": original_filename,
        "timestamp": timestamp,
        "raw_ocr_text": ocr_text,
        "raw_ocr_path": raw_txt_path,
        "extracted_workbook_path": extracted_xlsx_path,
        "extracted_workbook_name": os.path.basename(extracted_xlsx_path),
        "merged_workbook_path": merged_path,
        "debug_rows": debug_rows,
        "status_counts": status_counts,
        "total_fields": len(debug_rows),
        "ocr_reused": ocr_reused,
        "extracted": extracted,
        "ai_validation_results": ai_results,
        "ai_presence_results": ai_presence_results,
        "ai_validation_available": ai_validation_available,
        "ai_validation_detail": ai_availability.detail,
        "ai_technical_counts": ai_technical_counts,
        "ai_validation_counts": ai_validation_counts,
        "ai_presence_counts": ai_presence_counts,
        "correction_proposals": correction_proposals,
        "correction_available": correction_available,
        "correction_detail": correction_detail,
        "correction_technical_counts": correction_technical_counts,
        "correction_proposal_counts": correction_proposal_counts,
        "ai_mapping_results": ai_mapping_results,
        "ai_mapping_available": ai_mapping_available,
        "ai_mapping_detail": ai_mapping_availability.detail if ai_mapping_availability is not None else "No fields required a mapping attempt.",
        "ai_mapping_stage_state": ai_mapping_stage_state,
        "ai_mapping_stage_detail": ai_mapping_stage_detail,
        "ai_mapping_technical_counts": ai_mapping_technical_counts,
        "ai_mapping_status_counts": ai_mapping_status_counts,
        "source_line_pages": source_line_pages,
    }
