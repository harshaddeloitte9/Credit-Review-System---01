"""
Area 9 - Credit Memo: Summary.

Terminal aggregator - '9 - Credit Memo' has no calculations of its own; it
pulls a single status cell out of each of the other areas:
    B4  = '2-Financial Analysis'!B3    (Overall status: Approve/Reject)
    B8  = '3-Risk Scoring'!B10         (Overall status: risk grade text)
    B12 = '4-Covenants'!E14            (Overall Status: PASS/NOT PASS)
    B16 = '5-Collateral Assessment'!B12 (= D24, Collateral Rating text)
    B20 = '6-Limit Setting'!B40        (Conclusion: PASS/false)
    B24 = '7-Continuous Monitoring'!D10 (Overall Status: PASS/NOT PASS)
    B28 = '8-Portfolio Monitoring'!B4  (Overall status: Approve/Reject)

At stakeholder request, Continuous Monitoring is no longer shown as its
own standalone area (in the navigation or here) - its status is folded
into a "Covenant Assessment" grouping alongside financial and
non-financial covenant status, since app.py's Covenant Monitoring page
now presents all three together. B24's narrative (read by
load_credit_memo_commentary(), unchanged) is still sourced and shown -
just under that merged grouping instead of its own area. Non-financial
covenant status is always "Not Available / Requires Review" here, same
as on the Covenant Monitoring page - no field anywhere in this
application's data sources supports computing it.

Fund-Based / Non-Fund-Based Limit Setting (fb_result / nfb_result,
optional) are likewise displayed, not computed, here - they are already-
computed results app.py builds from the analyst-entered inputs on the
Limit Setting page (see limit_setting.py's build_fund_based_limit /
build_non_fund_based_limit) and simply passes in, the same way every
other *_result parameter below is a pre-built object, not raw data this
module derives anything from.

This module therefore does nothing but collect and format the already-
computed results of financial_analysis.py / risk_scoring.py (frozen) and
covenants.py / collateral_assessment.py / limit_setting.py /
continuous_monitoring.py / portfolio_monitoring.py (built earlier in this
same pass) - it re-derives none of them.

The reference workbook's A-column narrative paragraphs (A4, A8, A12, A16,
A20, A24, A28) are hand-written analyst prose - editorial judgement
("Key concern remains weak current ratio", "cyclical steel industry
exposure"), not a formula or anything derivable from the data. These are
read and shown as the reference workbook's own commentary, clearly
labelled as such, rather than regenerated - never fabricated as if this
application had produced its own written analysis.
"""

from dataclasses import dataclass

from openpyxl import load_workbook

CREDIT_MEMO_SHEET = "9 - Credit Memo"


@dataclass(frozen=True)
class CreditMemoArea:
    area_label: str
    status: object
    reference_commentary: str


@dataclass(frozen=True)
class CreditMemoResult:
    areas: list


_NARRATIVE_UNAVAILABLE = (
    "Reference workbook commentary — not dynamically generated. This workbook has no "
    "'9 - Credit Memo' narrative for this area."
)


def load_credit_memo_commentary(workbook_path) -> dict:
    """Returns a dict with all 7 keys always present; a value is the
    unavailable-placeholder string (never None, never Tata Steel's text)
    if the '9 - Credit Memo' sheet or that specific cell is missing from
    this workbook."""
    wb = load_workbook(workbook_path, data_only=True)
    if CREDIT_MEMO_SHEET not in wb.sheetnames:
        return {key: _NARRATIVE_UNAVAILABLE for key in (
            "financial_analysis", "risk_scoring", "covenants", "collateral_assessment",
            "limit_setting", "continuous_monitoring", "portfolio_monitoring",
        )}
    ws = wb[CREDIT_MEMO_SHEET]

    def _cell(coord):
        value = ws[coord].value
        return str(value) if value not in (None, "") else _NARRATIVE_UNAVAILABLE

    return {
        "financial_analysis": _cell("A4"),
        "risk_scoring": _cell("A8"),
        "covenants": _cell("A12"),
        "collateral_assessment": _cell("A16"),
        "limit_setting": _cell("A20"),
        "continuous_monitoring": _cell("A24"),
        "portfolio_monitoring": _cell("A28"),
    }


def _fmt_cr_or_review(value):
    return "Not Available / Requires Review" if value is None else f"₹{value:,.0f} Cr"


def build_credit_memo(fa, risk, covenants_result, collateral_result, limit_result,
                       monitoring_result, portfolio_result, commentary: dict,
                       fb_result=None, nfb_result=None) -> CreditMemoResult:
    """collateral_result / portfolio_result may be None when
    collateral_assessment.load_collateral_inputs / portfolio_monitoring.
    load_portfolio_inputs found no data for this workbook (see their own
    docstrings) - shown as "Not available" here rather than crashing or
    silently reusing a stale/other borrower's result. fb_result /
    nfb_result may be None (or carry None fields) when the analyst has
    not yet entered Fund-Based/Non-Fund-Based inputs on the Limit Setting
    page for this borrower - shown as "Not Available / Requires Review",
    never a fabricated figure."""
    limit_conclusion = limit_result.conclusion if isinstance(limit_result.conclusion, str) else "FALSE"
    collateral_status = collateral_result.rating if collateral_result else "Not available"
    portfolio_status = portfolio_result.overall_status if portfolio_result else "Not available"
    fb_status = _fmt_cr_or_review(fb_result.fb_requirement if fb_result else None)
    nfb_status = _fmt_cr_or_review(nfb_result.assessed_nfb_limit if nfb_result else None)

    # "Area N" labels below are a legacy numbering independent of the
    # app's current sidebar nav position (which has since gained a "Data
    # Review" step and renumbered further) - not the original reference
    # workbook's own sheet numbers either ('2-Financial Analysis',
    # '3-Risk Scoring', etc.), which stopped matching once Risk Scoring
    # moved to the end of the app's navigation as an "Additional
    # Feature". These labels are simply fixed, human-readable area names
    # at this point - kept as-is rather than renumbered again, since
    # several tests key off these exact strings and they carry no
    # behavioral meaning of their own.
    areas = [
        CreditMemoArea("Area 3 - Financial Analysis & Interpretation", fa.overall_status, commentary["financial_analysis"]),
        CreditMemoArea("Area 4 - Covenant Assessment — Financial Covenants", covenants_result.overall_status, commentary["covenants"]),
        CreditMemoArea("Area 4 - Covenant Assessment — Continuous / Early-Warning Monitoring", monitoring_result.overall_status, commentary["continuous_monitoring"]),
        CreditMemoArea("Area 4 - Covenant Assessment — Non-Financial Covenants", "Not Available / Requires Review",
                       "No field in the uploaded financial statement or the reference workbook captures "
                       "non-financial covenant compliance (borrowing restrictions, negative pledge, change "
                       "of control, asset disposal, regulatory/environmental compliance) - status is not "
                       "fabricated in its absence."),
        CreditMemoArea("Area 5 - Collateral Assessment", collateral_status, commentary["collateral_assessment"]),
        CreditMemoArea("Area 6 - Proposed Loan Impact & Limit Recommendation", limit_conclusion, commentary["limit_setting"]),
        CreditMemoArea("Area 6 - Limit Setting — Fund-Based (FB) Assessment", fb_status,
                       "Analyst-entered working-capital assessment (Limit Setting page) - no reference "
                       "workbook commentary applies."),
        CreditMemoArea("Area 6 - Limit Setting — Non-Fund-Based (NFB) Assessment", nfb_status,
                       "Analyst-entered LC/BG assessment (Limit Setting page) - no reference workbook "
                       "commentary applies."),
        CreditMemoArea("Area 7 - Portfolio Monitoring", portfolio_status, commentary["portfolio_monitoring"]),
        CreditMemoArea("Area 9 - Risk Scoring", risk.overall_status, commentary["risk_scoring"]),
    ]

    return CreditMemoResult(areas=areas)
