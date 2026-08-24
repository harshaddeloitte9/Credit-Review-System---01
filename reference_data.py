"""
Loads the Tata Steel reference dataset directly from the business/reference
workbook ("Credit Review Automation Example.xlsx") so the app can run the
current demo entirely offline - no PDF upload, no OCR.

Only cells that are genuinely raw inputs in '1-Input Data' (not Excel
formulas) are read here, plus the policy/assumption cells that '2-Financial
Analysis' and '3-Risk Scoring' already define (thresholds, weights,
qualitative scores, the risk-grade scale). Every formula that turns those
raw inputs into Financial Analysis / Risk Scoring metrics is re-implemented
in Python elsewhere (input_data_engine.py, financial_analysis.py,
risk_scoring.py) - nothing here is a derived or hardcoded number.
"""

import os
import re
from dataclasses import dataclass

from openpyxl import load_workbook

APP_DIRECTORY = os.path.dirname(os.path.abspath(__file__))

REFERENCE_WORKBOOK_FILE = "Credit Review Automation Example.xlsx"
REFERENCE_WORKBOOK_PATH = os.path.join(APP_DIRECTORY, REFERENCE_WORKBOOK_FILE)

INPUT_SHEET = "1-Input Data"
FINANCIAL_ANALYSIS_SHEET = "2-Financial Analysis"
RISK_SCORING_SHEET = "3-Risk Scoring"

_YEAR_RE = re.compile(r"(\d{4})")


class ReferenceDataError(RuntimeError):
    """Raised when the reference workbook is missing, or a cell isn't
    shaped the way the loader expects (e.g. a formula where a raw input
    was expected). Never guessed around - surfaced to the caller."""


@dataclass(frozen=True)
class YearFigure:
    current: float
    previous: float


@dataclass(frozen=True)
class ReferenceInputData:
    company_name: str
    statement_scope: str
    currency_unit: str
    fy_current_label: str
    fy_previous_label: str
    fy_current_year: int
    fy_previous_year: int

    # Profit & Loss (raw)
    revenue: YearFigure
    other_income: YearFigure
    total_expenses: YearFigure
    finance_costs: YearFigure
    depreciation: YearFigure
    interest_income: YearFigure
    exceptional_items: YearFigure
    total_tax_expense: YearFigure

    # Balance Sheet (raw)
    ppe: YearFigure
    total_assets: YearFigure
    total_equity: YearFigure
    borrowings_non_current: YearFigure
    borrowings_current: YearFigure
    lease_liabilities_non_current: YearFigure
    lease_liabilities_current: YearFigure
    current_investments: YearFigure
    cash: YearFigure
    other_bank_balances: YearFigure
    current_assets: YearFigure
    current_liabilities: YearFigure
    total_liabilities: YearFigure

    # Cash Flow (raw)
    operating_cash_flow: YearFigure
    income_taxes_paid: YearFigure
    interest_paid: YearFigure
    repayment_borrowings: YearFigure
    lease_payment: YearFigure
    capex: YearFigure
    dividend_paid: YearFigure
    total_comprehensive_income: YearFigure

    # Reported Ratios - Regulation 52(4) disclosures, "extracted as-is"
    # (these are inputs, not formulas, in the reference workbook)
    reported_dscr: YearFigure
    reported_iscr: YearFigure
    reported_current_ratio: YearFigure
    reported_debt_equity: YearFigure
    reported_debt_assets: YearFigure
    reported_ebitda_margin: YearFigure
    reported_net_profit_margin: YearFigure
    reported_net_worth: YearFigure

    # Same Regulation 52(4) disclosure block as the eight fields above
    # (rows 74 and 78, appended after the sheet's original row 73 - see
    # ocr_pipeline.py's ROW_MAPPING comment for why row 78, not 75).
    # Genuinely None (not a fabricated figure) for any document that
    # doesn't disclose these two ratios - not every filing includes them.
    reported_debtors_turnover_days: YearFigure
    reported_inventory_turnover_days: YearFigure

    # "Reported (validation)" column D - current year only, no FY-previous
    # equivalent exists in the sheet
    reported_total_income_current: float
    reported_pbt_before_exceptional_current: float
    reported_pbt_current: float
    reported_net_profit_current: float


RISK_COMPONENT_ROWS = {
    13: ("financial_strength", "Financial Strength"),
    14: ("industry_risk", "Industry Risk"),
    15: ("management_quality", "Management Quality"),
    16: ("leverage", "Leverage"),
    17: ("liquidity", "Liquidity"),
    18: ("debt_service_capacity", "Debt Service Capacity"),
    19: ("collateral_asset_coverage", "Collateral/Asset Coverage"),
    20: ("business_position", "Business Position"),
}

FINANCIAL_STRENGTH_SUBWEIGHT_ROWS = {
    30: "dscr",
    31: "iscr",
    32: "current_ratio",
    33: "debt_equity",
}


@dataclass(frozen=True)
class RiskScoringPolicy:
    component_weights: dict
    qualitative_scores: dict
    financial_strength_subweights: dict
    grade_scale: list


@dataclass(frozen=True)
class FinancialAnalysisPolicy:
    thresholds: dict


@dataclass(frozen=True)
class ReferenceBundle:
    input_data: ReferenceInputData
    risk_policy: RiskScoringPolicy
    fa_policy: FinancialAnalysisPolicy
    workbook_path: str


def _cell(ws, coord):
    # openpyxl's Worksheet has no real __contains__ (it falls back to
    # iterating rows), so "coord in ws" is silently always False and must
    # not be used here. ws[coord] itself never raises for a syntactically
    # valid reference - missing data is caught downstream via .value.
    return ws[coord]


def _num(ws, coord):
    cell = _cell(ws, coord)
    value = cell.value
    if value is None:
        return None
    if isinstance(value, str) and value.strip().startswith("="):
        raise ReferenceDataError(
            f"{ws.title}!{coord} was expected to be a raw input value but "
            f"contains a formula ({value!r}). Refusing to guess - the "
            f"reference workbook's structure may have changed."
        )
    if not isinstance(value, (int, float)):
        raise ReferenceDataError(
            f"{ws.title}!{coord} was expected to be a number but contains "
            f"{value!r}."
        )
    return float(value)


def _text(ws, coord):
    cell = _cell(ws, coord)
    value = cell.value
    if value is None:
        raise ReferenceDataError(f"{ws.title}!{coord} is empty.")
    return str(value).strip()


def _year_figure(ws, current_coord, previous_coord):
    return YearFigure(_num(ws, current_coord), _num(ws, previous_coord))


def _parse_header(a1_text):
    try:
        after_colon = a1_text.split(":", 1)[1].strip()
        name_scope, unit_part = after_colon.rsplit("(", 1)
        company, scope = [p.strip() for p in name_scope.strip().rstrip(",").split(",")]
        unit = unit_part.rstrip(")").strip()
        return company, scope, unit
    except Exception as exc:
        raise ReferenceDataError(
            f"Could not parse company/scope/unit out of 1-Input Data!A1 "
            f"({a1_text!r}): {exc}"
        )


def _parse_year(label):
    match = _YEAR_RE.search(label)
    if not match:
        raise ReferenceDataError(f"Could not find a 4-digit year in {label!r}.")
    return int(match.group(1))


def _load_input_data(ws) -> ReferenceInputData:
    company, scope, unit = _parse_header(_text(ws, "A1"))
    fy_current_label = _text(ws, "B4")
    fy_previous_label = _text(ws, "C4")

    return ReferenceInputData(
        company_name=company,
        statement_scope=scope,
        currency_unit=unit,
        fy_current_label=fy_current_label,
        fy_previous_label=fy_previous_label,
        fy_current_year=_parse_year(fy_current_label),
        fy_previous_year=_parse_year(fy_previous_label),

        revenue=_year_figure(ws, "B5", "C5"),
        other_income=_year_figure(ws, "B6", "C6"),
        total_expenses=_year_figure(ws, "B8", "C8"),
        finance_costs=_year_figure(ws, "B9", "C9"),
        depreciation=_year_figure(ws, "B10", "C10"),
        interest_income=_year_figure(ws, "B11", "C11"),
        exceptional_items=_year_figure(ws, "B13", "C13"),
        total_tax_expense=_year_figure(ws, "B15", "C15"),

        ppe=_year_figure(ws, "B23", "C23"),
        total_assets=_year_figure(ws, "B24", "C24"),
        total_equity=_year_figure(ws, "B25", "C25"),
        borrowings_non_current=_year_figure(ws, "B26", "C26"),
        borrowings_current=_year_figure(ws, "B27", "C27"),
        lease_liabilities_non_current=_year_figure(ws, "B28", "C28"),
        lease_liabilities_current=_year_figure(ws, "B29", "C29"),
        current_investments=_year_figure(ws, "B31", "C31"),
        cash=_year_figure(ws, "B32", "C32"),
        other_bank_balances=_year_figure(ws, "B33", "C33"),
        current_assets=_year_figure(ws, "B35", "C35"),
        current_liabilities=_year_figure(ws, "B36", "C36"),
        total_liabilities=_year_figure(ws, "B37", "C37"),

        operating_cash_flow=_year_figure(ws, "B40", "C40"),
        income_taxes_paid=_year_figure(ws, "B41", "C41"),
        interest_paid=_year_figure(ws, "B42", "C42"),
        repayment_borrowings=_year_figure(ws, "B43", "C43"),
        lease_payment=_year_figure(ws, "B44", "C44"),
        capex=_year_figure(ws, "B45", "C45"),
        dividend_paid=_year_figure(ws, "B46", "C46"),
        total_comprehensive_income=_year_figure(ws, "B47", "C47"),

        reported_dscr=_year_figure(ws, "B49", "D49"),
        reported_iscr=_year_figure(ws, "B50", "D50"),
        reported_current_ratio=_year_figure(ws, "B51", "D51"),
        reported_debt_equity=_year_figure(ws, "B52", "D52"),
        reported_debt_assets=_year_figure(ws, "B53", "D53"),
        reported_ebitda_margin=_year_figure(ws, "B54", "D54"),
        reported_net_profit_margin=_year_figure(ws, "B55", "D55"),
        reported_net_worth=_year_figure(ws, "B56", "D56"),
        reported_debtors_turnover_days=_year_figure(ws, "B74", "D74"),
        reported_inventory_turnover_days=_year_figure(ws, "B78", "D78"),

        reported_total_income_current=_num(ws, "D7"),
        reported_pbt_before_exceptional_current=_num(ws, "D12"),
        reported_pbt_current=_num(ws, "D14"),
        reported_net_profit_current=_num(ws, "D16"),
    )


def _load_financial_analysis_policy(ws) -> FinancialAnalysisPolicy:
    thresholds = {
        "revenue": _num(ws, "G7"),
        "ebitda_margin_growth": _num(ws, "G8"),
        "net_profit": _num(ws, "G9"),
        "total_assets": _num(ws, "G10"),
        "net_worth": _num(ws, "G11"),
        "operating_cash_flow": _num(ws, "G12"),
        "ocf_to_debt": _num(ws, "G13"),
        "total_debt": _num(ws, "G14"),
        "debt_equity": _num(ws, "G18"),
        "dscr": _num(ws, "G19"),
        "iscr": _num(ws, "G20"),
        "current_ratio": _num(ws, "G21"),
        "debt_assets": _num(ws, "G22"),
        "debtors_days": _num(ws, "G23"),
        "inventory_days": _num(ws, "G24"),
        "ebitda_margin_pct": _num(ws, "G25"),
        "net_profit_margin_pct": _num(ws, "G26"),
    }
    return FinancialAnalysisPolicy(thresholds=thresholds)


def _load_risk_policy(ws) -> RiskScoringPolicy:
    component_weights = {}
    qualitative_scores = {}

    for row, (key, expected_label) in RISK_COMPONENT_ROWS.items():
        label = _text(ws, f"A{row}")
        if label != expected_label:
            raise ReferenceDataError(
                f"{RISK_SCORING_SHEET}!A{row} was expected to be "
                f"{expected_label!r} but is {label!r} - row layout may "
                f"have changed."
            )
        component_weights[key] = _num(ws, f"B{row}")
        if key != "financial_strength":
            # C13 (Financial Strength) is itself a formula in the reference
            # workbook (=F35) - that is the one component this app computes,
            # not a policy input, so it is deliberately not read here.
            qualitative_scores[key] = _num(ws, f"C{row}")

    financial_strength_subweights = {
        key: _num(ws, f"E{row}")
        for row, key in FINANCIAL_STRENGTH_SUBWEIGHT_ROWS.items()
    }

    grade_scale = []
    for row in range(4, 9):
        grade_scale.append(
            {
                "score": int(_num(ws, f"A{row}")),
                "definition": _text(ws, f"B{row}"),
                "risk_interpretation": _text(ws, f"C{row}"),
            }
        )

    return RiskScoringPolicy(
        component_weights=component_weights,
        qualitative_scores=qualitative_scores,
        financial_strength_subweights=financial_strength_subweights,
        grade_scale=grade_scale,
    )


def load_reference_bundle(path: str = REFERENCE_WORKBOOK_PATH) -> ReferenceBundle:
    if not os.path.exists(path):
        raise ReferenceDataError(
            f"Reference workbook not found: {path}. Make sure "
            f"'{REFERENCE_WORKBOOK_FILE}' is present in the app folder."
        )

    wb = load_workbook(path, data_only=False)

    for sheet in (INPUT_SHEET, FINANCIAL_ANALYSIS_SHEET, RISK_SCORING_SHEET):
        if sheet not in wb.sheetnames:
            raise ReferenceDataError(
                f"Reference workbook is missing the '{sheet}' sheet."
            )

    input_data = _load_input_data(wb[INPUT_SHEET])
    fa_policy = _load_financial_analysis_policy(wb[FINANCIAL_ANALYSIS_SHEET])
    risk_policy = _load_risk_policy(wb[RISK_SCORING_SHEET])

    return ReferenceBundle(
        input_data=input_data,
        risk_policy=risk_policy,
        fa_policy=fa_policy,
        workbook_path=path,
    )
