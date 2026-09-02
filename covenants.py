"""
Area 4 - Covenant Monitoring.

Reference workbook dependency (verified from actual formulas, not assumed):
    '4-Covenants'!C11 = '2-Financial Analysis'!C19   (DSCR, reported basis)
    '4-Covenants'!C12 = '2-Financial Analysis'!C20   (ISCR, reported basis)

So this module consumes financial_analysis.py's ALREADY-COMPUTED output
(the frozen module) - it does not re-read Input Data and does not
recompute DSCR/ISCR itself.

Policy inputs (the reference workbook's own "Yellow fill = Assumption /
policy parameter" cells B11, B12) are read live from the reference
workbook rather than hardcoded as Python literals, same pattern as
reference_data.py uses for the frozen modules.

Formulas reproduced exactly as found in '4-Covenants':
    Headroom  = (Actual - Threshold) / Threshold                  (D11, D12)
    RAG       = GREEN if Headroom > 20%, AMBER if >= 5%, else RED  (E11, E12)
    Overall   = "PASS" unless any covenant is AMBER or RED         (E14)

FIXED (not a reference-workbook formula, a gap this module's own Python
introduced): the reference workbook's E14 has no notion of a covenant
being un-evaluable at all (every cell in its worked example is filled
in), so it offers no precedent either way for what "PASS unless AMBER/RED"
should do when a covenant's underlying figure is missing. The original
Python here interpreted that silence as "exclude it and pass on whatever
is left" - so a document missing ISCR (not every filing discloses it) but
passing DSCR would show a clean overall "PASS", never having actually
confirmed ISCR compliance. build_covenants() below now treats a missing
REQUIRED covenant as blocking a clean PASS (see its own comment) - a
breach still takes precedence over a gap, but a gap can no longer read as
compliance.
"""

from dataclasses import dataclass, field
from typing import Optional

from openpyxl import load_workbook

COVENANTS_SHEET = "4-Covenants"
NOT_AVAILABLE = "Not Available / Requires Review"


@dataclass(frozen=True)
class CovenantPolicy:
    dscr_minimum: float
    iscr_minimum: float


@dataclass(frozen=True)
class CovenantRow:
    label: str
    threshold: float
    actual: Optional[float]
    headroom: Optional[float]
    rag: str


@dataclass(frozen=True)
class CovenantsResult:
    rows: list
    overall_status: str
    overall_status_note: str = field(default="")


def load_covenant_policy(workbook_path) -> CovenantPolicy:
    ws = load_workbook(workbook_path, data_only=True)[COVENANTS_SHEET]
    return CovenantPolicy(
        dscr_minimum=float(ws["B11"].value),
        iscr_minimum=float(ws["B12"].value),
    )


def _rag(headroom):
    # '4-Covenants'!E11/E12: =IF(D>0.2,"GREEN",IF(D>=0.05,"AMBER","RED"))
    if headroom > 0.2:
        return "GREEN"
    if headroom >= 0.05:
        return "AMBER"
    return "RED"


def _covenant_row(label, threshold, actual):
    # DSCR/ISCR (reported basis) are this document's "Reported Ratios"
    # disclosure - not every filing includes it (see
    # financial_analysis.py's identical guard). A covenant that cannot be
    # evaluated shows as such rather than crashing on (None - threshold).
    if actual is None:
        return CovenantRow(label, threshold, None, None, NOT_AVAILABLE)
    headroom = (actual - threshold) / threshold
    return CovenantRow(label, threshold, actual, headroom, _rag(headroom))


def build_covenants(fa, policy: CovenantPolicy) -> CovenantsResult:
    rows_by_key = {r.key: r for r in fa.ratio_rows}
    dscr_actual = rows_by_key["dscr"].current
    iscr_actual = rows_by_key["iscr"].current

    rows = [
        _covenant_row("DSCR, reported basis (minimum)", policy.dscr_minimum, dscr_actual),
        _covenant_row("ISCR, reported basis (minimum)", policy.iscr_minimum, iscr_actual),
    ]

    evaluable_rows = [r for r in rows if r.rag != NOT_AVAILABLE]
    unavailable_labels = [r.label.split(",")[0] for r in rows if r.rag == NOT_AVAILABLE]
    # '4-Covenants'!E14: PASS unless any (evaluable) covenant is AMBER or RED
    breach = any(r.rag in ("AMBER", "RED") for r in evaluable_rows)

    # Explicit precedence across the three states this gate can land in -
    # a breach is reported even if another covenant is ALSO missing (the
    # breach is the more urgent fact); short of a breach, any required
    # covenant that could not be evaluated must stop this from reading as
    # a clean PASS. Fixed here: previously, a missing ISCR alongside a
    # passing DSCR still produced an overall "PASS" - compliance was never
    # actually confirmed in full, just silently treated as if it had been.
    if breach:
        overall_status = "NOT PASS"
        note = (
            f"{' and '.join(unavailable_labels)} could not be evaluated (not disclosed in this "
            f"document) and is excluded from this status."
            if unavailable_labels else ""
        )
    elif not evaluable_rows:
        overall_status = NOT_AVAILABLE
        note = "Neither covenant could be evaluated - this document's Reported Ratios disclosure includes neither DSCR nor ISCR."
    elif unavailable_labels:
        overall_status = NOT_AVAILABLE
        note = (
            f"{' and '.join(unavailable_labels)} could not be evaluated (not disclosed in this "
            f"document). The covenant(s) that could be checked are within limits, but overall "
            f"compliance cannot be confirmed without {' and '.join(unavailable_labels)} - not "
            f"assumed compliant in its absence."
        )
    else:
        overall_status = "PASS"
        note = ""

    return CovenantsResult(rows=rows, overall_status=overall_status, overall_status_note=note)
