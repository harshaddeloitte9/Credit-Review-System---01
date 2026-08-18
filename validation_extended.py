"""
Validation for modules 4-9 (Covenants, Collateral Assessment, Limit Setting,
Continuous Monitoring, Portfolio Monitoring, Credit Memo) against the
reference workbook's own cached cell values.

Deliberately kept separate from validation.py, which validates the frozen
1-3 modules only. Nothing in that file is imported, wrapped, or modified
here - this module only reads the reference workbook (data_only=True) and
compares against the new modules' own computed output, using the same
ValidationRow shape and comparison semantics as validation.py for
consistency, without touching it.
"""

from dataclasses import dataclass
from typing import Optional

from openpyxl import load_workbook


@dataclass(frozen=True)
class ValidationRow:
    metric: str
    calculated: object
    reference: object
    difference: Optional[object]
    status: str


def _compare(metric, calculated, reference, tolerance):
    if calculated is None or reference is None:
        return ValidationRow(metric, calculated, reference, None, "NOT COMPARABLE")
    if isinstance(calculated, str) or isinstance(reference, str):
        status = "MATCH" if str(calculated).strip() == str(reference).strip() else "DISCREPANCY"
        return ValidationRow(metric, calculated, reference, None, status)
    diff = calculated - reference
    status = "MATCH" if abs(diff) <= tolerance else "DISCREPANCY"
    return ValidationRow(metric, calculated, reference, diff, status)


def validate_covenants(result, workbook_path):
    ws = load_workbook(workbook_path, data_only=True)["4-Covenants"]
    rows = [
        _compare("DSCR headroom (D11)", result.rows[0].headroom, ws["D11"].value, 1e-6),
        _compare("DSCR RAG (E11)", result.rows[0].rag, ws["E11"].value, 0),
        _compare("ISCR headroom (D12)", result.rows[1].headroom, ws["D12"].value, 1e-6),
        _compare("ISCR RAG (E12)", result.rows[1].rag, ws["E12"].value, 0),
        _compare("Overall status (E14)", result.overall_status, ws["E14"].value, 0),
    ]
    return rows


def validate_limit_setting(result, workbook_path):
    ws = load_workbook(workbook_path, data_only=True)["6-Limit Setting"]
    rows = [
        _compare("Annual instalment (B7)", result.annual_instalment, ws["B7"].value, 1e-4),
    ]
    for year_row, cell_row in zip(result.schedule, range(11, 16)):
        rows.append(_compare(f"Year {year_row.year} opening balance (B{cell_row})", year_row.opening_balance, ws[f"B{cell_row}"].value, 1e-4))
        rows.append(_compare(f"Year {year_row.year} interest (C{cell_row})", year_row.interest, ws[f"C{cell_row}"].value, 1e-4))
        rows.append(_compare(f"Year {year_row.year} principal (D{cell_row})", year_row.principal, ws[f"D{cell_row}"].value, 1e-4))
        rows.append(_compare(f"Year {year_row.year} closing balance (E{cell_row})", year_row.closing_balance, ws[f"E{cell_row}"].value, 1e-3))
    rows.append(_compare("Implied existing debt service (B19)", result.implied_existing_debt_service, ws["B19"].value, 1e-4))
    rows.append(_compare("Total pro-forma debt service (B21)", result.total_proforma_debt_service, ws["B21"].value, 1e-4))
    rows.append(_compare("EBIT (B22)", result.ebit, ws["B22"].value, 1e-4))
    rows.append(_compare("Pro-forma DSCR (B23)", result.proforma_dscr, ws["B23"].value, 1e-6))
    rows.append(_compare("Pro-forma DSCR status (B25)", result.proforma_dscr_status, ws["B25"].value, 0))
    rows.append(_compare("Max serviceable debt service (B31)", result.max_serviceable_debt_service, ws["B31"].value, 1e-4))
    rows.append(_compare("Available additional debt service (B33)", result.available_additional_debt_service, ws["B33"].value, 1e-4))
    rows.append(_compare("Capacity cap (B34)", result.capacity_cap, ws["B34"].value, 1e-3))
    rows.append(_compare("Recommended limit (B39)", result.recommended_limit, ws["B39"].value, 1e-4))
    rows.append(_compare("Conclusion (B40)", result.conclusion, ws["B40"].value, 0))
    return rows


def validate_continuous_monitoring(result, workbook_path):
    ws = load_workbook(workbook_path, data_only=True)["7-Continuous Monitoring"]
    trigger_by_label = {t.label: t for t in result.triggers}
    rows = [
        _compare("Revenue decline YoY status (D4)", trigger_by_label["Revenue decline YoY"].status, ws["D4"].value, 0),
        _compare("Total debt increase YoY status (D5)", trigger_by_label["Total debt increase YoY"].status, ws["D5"].value, 0),
        _compare("Reported DSCR deterioration YoY status (D6)", trigger_by_label["Reported DSCR deterioration YoY"].status, ws["D6"].value, 0),
        _compare("Any covenant Amber/Red status (D7)", trigger_by_label["Any covenant on Amber/Red"].status, ws["D7"].value, 0),
        _compare("Overall status (D10)", result.overall_status, ws["D10"].value, 0),
    ]
    # D8/D9 (external rating, adverse news) are intentionally shown as "not
    # monitored" rather than asserting the reference workbook's static
    # "PASS" snapshot as a live fact - see continuous_monitoring.py's
    # module docstring. Reported here as not comparable by design, not as
    # a discrepancy.
    rows.append(ValidationRow(
        "External rating / adverse news (D8, D9)", "NOT MONITORED (intentional)",
        f"{ws['D8'].value} / {ws['D9'].value}", None,
        "NOT COMPARABLE - intentionally not asserted, see module docstring",
    ))
    return rows


def validate_credit_memo(result, workbook_path):
    ws = load_workbook(workbook_path, data_only=True)["9 - Credit Memo"]
    cells = ["B4", "B8", "B12", "B16", "B20", "B24", "B28"]
    rows = [
        _compare(f"{area.area_label} status ({cell})", area.status, ws[cell].value, 0)
        for area, cell in zip(result.areas, cells)
    ]
    return rows


def validate_portfolio_monitoring(result, workbook_path):
    ws = load_workbook(workbook_path, data_only=True)["8-Portfolio Monitoring"]
    rows = [
        _compare("Average exposure per account (C10)", result.average_exposure_per_account, ws["C10"].value, 1e-4),
        _compare("Borrower concentration % (C12)", result.borrower_concentration_pct, ws["C12"].value, 1e-6),
        _compare("Gross NPL ratio % (C14)", result.npl_ratio_pct, ws["C14"].value, 1e-6),
        _compare("Portfolio exposure post sanction (C15)", result.portfolio_exposure_post_sanction, ws["C15"].value, 1e-4),
        _compare("Single borrower exposure % (C16)", result.single_borrower_exposure_pct, ws["C16"].value, 1e-6),
        _compare("NPL risk band (C20)", result.criteria[0].band, ws["C20"].value, 0),
        _compare("Borrower concentration risk band (C21)", result.criteria[1].band, ws["C21"].value, 0),
        _compare("Manufacturing sector exposure band (C22)", result.criteria[2].band, ws["C22"].value, 0),
        _compare("Overall status (B4)", result.overall_status, ws["B4"].value, 0),
    ]
    return rows


def validate_collateral_assessment(result, workbook_path):
    ws = load_workbook(workbook_path, data_only=True)["5-Collateral Assessment"]
    rows = []
    for asset, row in zip(result.assets, (15, 16, 17, 18, 19)):
        rows.append(_compare(f"{asset.label} adjusted value (D{row})", asset.adjusted_value, ws[f"D{row}"].value, 1e-4))
    rows.append(_compare("Total adjusted collateral value (D21)", result.total_adjusted_value, ws["D21"].value, 1e-4))
    rows.append(_compare("Debt exposure (D22)", result.debt_exposure, ws["D22"].value, 1e-4))
    rows.append(_compare("Collateral coverage ratio (D23)", result.coverage_ratio, ws["D23"].value, 1e-6))
    rows.append(_compare("Collateral rating (D24)", result.rating, ws["D24"].value, 0))
    return rows
